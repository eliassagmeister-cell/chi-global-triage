#!/usr/bin/env python3
# CHI Global Triage: INFORM Severity x FTS funding.
# Pulls INFORM Severity (HDX) + FTS funding (OCHA) server-side, joins on ISO3,
# computes funding received per person in need, writes index.html and serves it.
#   python3 main.py         build + serve  (Replit "Run")
#   python3 main.py build   build only     (Scheduled Deployment)
import json, os, io, sys, datetime, urllib.request

FTS_YEAR = 2026
HDX_PKG  = "inform-global-crisis-severity-index"
UA = {"User-Agent": "CHI-triage/1.0"}
HERE = os.path.dirname(os.path.abspath(__file__))

def _get(url, timeout=90):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
        return r.read()

def load_fallback():
    try:
        return json.load(open(os.path.join(HERE, "data_fallback.json"), encoding="utf-8"))
    except Exception:
        return []

def parse_inform(wb):
    ac = wb["INFORM Severity - all crises"]
    rows = list(ac.iter_rows(values_only=True))
    hi = None; hdr = None
    for i, r in enumerate(rows[:8]):
        cells = [str(c).strip() if c else "" for c in (r or [])]
        if "CRISIS" in cells and "ISO3" in cells:
            hi, hdr = i, cells; break
    if hi is None:
        raise ValueError("all-crises header not found")
    def col(*names):
        for n in names:
            if n in hdr: return hdr.index(n)
        return None
    i_c = col("CRISIS"); i_cid = col("CRISIS ID")
    i_co = col("COUNTRY"); i_iso = col("ISO3"); i_dr = col("DRIVERS")
    sev_idx = [j for j, h in enumerate(hdr) if h.startswith("INFORM Severity")]
    i_sev = sev_idx[0]; i_cat = sev_idx[2] if len(sev_idx) >= 3 else None
    i_tr = col("Trend (last 3 months)", "Trend"); i_rel = col("Reliability")
    def g(r, i): return r[i] if (i is not None and i < len(r)) else None
    base = []
    for r in rows[hi + 1:]:
        if not r: continue
        co = g(r, i_co)
        try: sev = float(g(r, i_sev))
        except (TypeError, ValueError): continue
        if not co: continue
        base.append({"c": str(g(r, i_c) or co), "cid": str(g(r, i_cid) or ""),
                     "co": str(co), "iso": str(g(r, i_iso) or ""), "dr": str(g(r, i_dr) or ""),
                     "sev": round(sev, 1), "cat": str(g(r, i_cat) or ""),
                     "tr": str(g(r, i_tr) or ""), "rel": str(g(r, i_rel) or ""), "pin": None})
    pin = {}
    if "Conditions of people affected" in wb.sheetnames:
        cr = list(wb["Conditions of people affected"].iter_rows(values_only=True))
        def num(x):
            try: return float(x)
            except (TypeError, ValueError): return 0.0
        l3, l4, l5, idc = 7, 8, 9, 2
        for r in cr[:8]:
            cells = [str(c) if c is not None else "" for c in (r or [])]
            for j, c in enumerate(cells):
                cl = c.lower()
                if "# of people" in cl and "moderat" in cl: l3 = j
                elif "# of people" in cl and "severe" in cl: l4 = j
                elif "# of people" in cl and "extreme" in cl: l5 = j
            if "CRISIS ID" in cells: idc = cells.index("CRISIS ID")
        for r in cr:
            cid = r[idc] if idc < len(r) else None
            if not cid or not str(cid).strip(): continue
            v = num(r[l3] if l3 < len(r) else 0) + num(r[l4] if l4 < len(r) else 0) + num(r[l5] if l5 < len(r) else 0)
            if v > 0: pin[str(cid)] = int(round(v * 1e6))
    for b in base:
        b["pin"] = pin.get(b["cid"]); b.pop("cid", None)
    if len(base) < 20:
        raise ValueError("parsed too few crises")
    return base

def fetch_inform():
    try:
        import openpyxl, io
        meta = json.loads(_get("https://data.humdata.org/api/3/action/package_show?id=" + HDX_PKG))
        res = meta["result"]["resources"]
        pick = None
        for x in res:
            if str(x.get("format", "")).lower() == "xlsx":
                pick = x
                if "severity" in (str(x.get("name", "")) + str(x.get("url", ""))).lower():
                    break
        if not pick:
            raise ValueError("no xlsx on HDX")
        raw = _get(pick.get("download_url") or pick.get("url"))
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        return parse_inform(wb), "live from HDX"
    except Exception as e:
        return load_fallback(), "fallback snapshot (%s)" % str(e)[:70]

def fetch_fts(year):
    try:
        pj = json.loads(_get("https://api.hpc.tools/v1/public/plan/year/%d" % year))
        fj = json.loads(_get("https://api.hpc.tools/v1/public/fts/flow?year=%d&groupby=plan" % year))
        fund_by_plan = {}
        try:
            objs = fj["data"]["report3"]["fundingTotals"]["objects"]
        except (KeyError, TypeError):
            objs = []
        for o in objs:
            for sfo in (o.get("singleFundingObjects") or []):
                pid = sfo.get("id")
                if pid is not None:
                    fund_by_plan[pid] = float(sfo.get("totalFunding") or 0)
        m = {}
        for p in (pj.get("data") or []):
            pid = p.get("id")
            req = float(p.get("revisedRequirements") or p.get("origRequirements") or 0)
            fund = fund_by_plan.get(pid, 0.0)
            isos = [l.get("iso3") for l in (p.get("locations") or []) if l.get("iso3")]
            isos = list(dict.fromkeys(isos))
            if not isos:
                continue
            n = len(isos)
            for iso in isos:
                d = m.setdefault(iso, {"req": 0.0, "fund": 0.0})
                d["req"] += req / n
                d["fund"] += fund / n
        if not m:
            raise ValueError("no plans")
        return m, "live %d via hpc.tools" % year
    except Exception as e:
        return {}, "unavailable (%s)" % str(e)[:70]

def enrich(crises, fts):
    pin_by_iso = {}
    for d in crises:
        if d.get("iso") and d.get("pin"):
            pin_by_iso[d["iso"]] = pin_by_iso.get(d["iso"], 0) + d["pin"]
    for d in crises:
        f = fts.get(d.get("iso")) if d.get("iso") else None
        d["cov"] = round(100 * f["fund"] / f["req"]) if (f and f.get("req")) else None
        tot = pin_by_iso.get(d.get("iso"), 0)
        d["fpp"] = round(f["fund"] / tot) if (f and tot) else None
    return crises


HTML = r'''<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Severity and Neglectedness of Humanitarian Crises</title><style>
:root{--bg:#fafaf9;--card:#fff;--ink:#1c1917;--mut:#78716c;--line:#e7e5e4;--accent:#0f766e}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif}
.wrap{max-width:1000px;margin:0 auto;padding:26px 20px 60px}
h1{font-size:21px;margin:0 0 3px}a{color:var(--accent)}
.prov{color:var(--mut);font-size:12px;margin-bottom:16px}
.about{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin-bottom:14px;font-size:13px;line-height:1.55}
.about p{margin:0 0 10px}
.deft{font-size:11px;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);margin:10px 0 5px;font-weight:600}
.about ul{margin:0;padding-left:18px}.about li{margin:3px 0;color:var(--ink)}.about li b{color:var(--ink)}
.ctrls{display:flex;gap:20px;flex-wrap:wrap;align-items:center;background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px;margin-bottom:14px}
.ctl{flex:1;min-width:240px}.lab{font-size:12px;color:var(--mut);margin-bottom:6px}.lab b{color:var(--ink)}
input[type=range]{width:100%;accent-color:var(--accent)}
.wts{font-size:12px;color:var(--mut)}.wts b{color:var(--ink)}
table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--line);border-radius:10px;overflow:hidden}
th,td{text-align:left;padding:8px 10px;border-bottom:1px solid var(--line);font-size:13px}
th{font-size:11px;color:var(--mut);text-transform:uppercase;letter-spacing:.03em;background:#fbfbfa;white-space:nowrap}
th[data-k]{cursor:pointer;user-select:none}th[data-k]:hover{color:var(--ink)}
th.r,td.r{text-align:right}tr:last-child td{border-bottom:0}
.rk{color:var(--mut);text-align:right}.nm{font-weight:600}.iso{color:var(--mut);font-size:11px;margin-left:6px}
.pill{display:inline-block;min-width:32px;text-align:center;color:#fff;border-radius:6px;padding:2px 6px;font-weight:700;font-size:12px}
.gap{color:#b45309;font-weight:600}
.impc{min-width:140px}
.impc .barwrap{position:relative;background:#eef0ef;border-radius:5px;height:19px;overflow:hidden}
.impc .bar{height:100%;background:var(--accent)}
.impc .scv{position:absolute;right:7px;top:0;line-height:19px;font-size:11px;font-weight:700;color:var(--ink)}
.below{margin-top:12px;font-size:11.5px;color:var(--mut);line-height:1.6}
.below .flag{color:#9a3412}
</style></head><body><div class="wrap">
<h1>Severity and Neglectedness of Humanitarian Crises</h1>
<div class="prov" id="prov"></div>
<div class="about">
<p>Every active humanitarian crisis, ranked by a score that blends how severe it is with how underfunded it is. Move the &alpha; slider to weight neglect against severity: at 0 the ranking is pure severity, at 1 it is pure funding gap.</p>
<div class="deft">Sources</div>
<ul>
<li><b>INFORM Severity</b> (ACAPS) &mdash; a 0&ndash;10 composite of a crisis's impact, the conditions of the people affected, and its complexity, refreshed monthly.</li>
<li><b>OCHA FTS</b> &mdash; humanitarian funding reported against each crisis's UN-coordinated appeal, measured against what that appeal requires.</li>
</ul>
<div class="deft">What each column means</div>
<ul>
<li><b>Severity</b> &mdash; the INFORM Severity score, 0&ndash;10.</li>
<li><b>Funding gap</b> &mdash; the share of the appeal still unfunded (1 &minus; coverage); higher means more neglected.</li>
<li><b>$ / person</b> &mdash; funding received divided by the country's total people in need.</li>
<li><b>Impact opportunity</b> &mdash; the ranking score, &alpha; &times; funding gap + (1 &minus; &alpha;) &times; severity, shown out of 100.</li>
<li><b>&alpha;</b> &mdash; the slider below: the weight given to the funding gap (neglect) versus severity.</li>
</ul>
</div>
<div class="ctrls">
<div class="ctl"><div class="lab">&alpha; (funding-gap weight) = <b id="av">0.35</b></div><input id="a" type="range" min="0" max="1" step="0.05" value="0.35"></div>
<div class="wts">severity <b id="sw">65</b>% &middot; funding gap <b id="fw">35</b>%</div>
</div>
<table><thead><tr><th class="r">#</th><th data-k="co">Crisis</th><th class="r" data-k="sev">Severity</th><th class="r" data-k="gap">Funding gap</th><th class="r" data-k="fpp">$ / person</th><th data-k="imp">Impact opportunity (0&ndash;100)</th></tr></thead><tbody id="tb"></tbody></table>
<div class="below">
<div id="noapp"></div>
<div>Ranks the <b id="nn"></b> crises with an FTS appeal this year; regional (multi-country) plans are split evenly across their countries. Click any column header to sort by it. Data live from INFORM Severity (ACAPS/HDX) and OCHA FTS.</div>
</div>
</div>
<script>
var D=__DATA__,M=__META__;
var CAT={"Very High":"#b91c1c","High":"#ea580c","Medium":"#d97706","Low":"#65a30d","Very Low":"#4d7c0f"};
var st={k:"imp",dir:-1,A:0.35};
function render(){
  st.A=+document.getElementById('a').value;
  document.getElementById('av').textContent=st.A.toFixed(2);
  document.getElementById('sw').textContent=Math.round((1-st.A)*100);
  document.getElementById('fw').textContent=Math.round(st.A*100);
  D.forEach(function(d){if(d.cov!=null){var sf=d.sev/10,ff=1-d.cov/100;d.imp=st.A*ff+(1-st.A)*sf;d.gap=100-d.cov;}});
  var rows=D.filter(function(d){return d.cov!=null;});
  rows.sort(function(a,b){var A=a[st.k],B=b[st.k];if(A==null)return 1;if(B==null)return -1;if(typeof A==="string")return st.dir*A.localeCompare(B);return st.dir*(A-B);});
  document.getElementById('tb').innerHTML=rows.map(function(d,i){var c=CAT[d.cat]||"#78716c";var sc=Math.round(d.imp*100);
    return '<tr><td class="rk">'+(i+1)+'</td><td><span class="nm">'+d.co+'</span><span class="iso">'+(d.iso||'')+'</span></td>'+
    '<td class="r"><span class="pill" style="background:'+c+'">'+d.sev.toFixed(1)+'</span></td>'+
    '<td class="r gap">'+d.gap+'%</td>'+
    '<td class="r">'+(d.fpp==null?'&ndash;':'$'+(+d.fpp).toLocaleString())+'</td>'+
    '<td class="impc"><div class="barwrap"><div class="bar" style="width:'+sc+'%"></div><span class="scv">'+sc+'</span></div></td></tr>';}).join('');
  document.getElementById('nn').textContent=rows.length;
  var na=D.filter(function(d){return d.cov==null && d.sev>=7;}).sort(function(a,b){return b.sev-a.sev;});
  document.getElementById('noapp').innerHTML=na.length?('<span class="flag">Note:</span> some severe crises have no FTS appeal and so can\'t be scored on funding &mdash; worth a manual look: '+na.map(function(d){return d.co+' '+d.sev.toFixed(1);}).join(', ')+'.'):'';
}
document.getElementById('a').addEventListener('input',render);
[].slice.call(document.querySelectorAll('th[data-k]')).forEach(function(th){th.onclick=function(){var k=th.dataset.k;if(st.k===k)st.dir*=-1;else{st.k=k;st.dir=(k==="co")?1:-1;}render();};});
document.getElementById('prov').textContent='Generated '+M.date+' — INFORM '+M.inform+', FTS '+M.fts;
render();
</script></body></html>'''

def build():
    crises, isrc = fetch_inform()
    fts, fsrc = fetch_fts(FTS_YEAR)
    enrich(crises, fts)
    crises.sort(key=lambda d: -d["sev"])
    meta = {"date": datetime.date.today().isoformat(), "inform": isrc, "fts": fsrc}
    html = HTML.replace("__DATA__", json.dumps(crises, ensure_ascii=False)).replace("__META__", json.dumps(meta, ensure_ascii=False))
    open(os.path.join(HERE, "index.html"), "w", encoding="utf-8").write(html)
    print("built:", len(crises), "crises | INFORM:", isrc, "| FTS:", fsrc)
    return meta

import time, threading
REBUILD_DAYS = float(os.environ.get("REBUILD_DAYS", "25"))
_lock = threading.Lock()

def _stale():
    p = os.path.join(HERE, "index.html")
    if not os.path.exists(p):
        return True
    return (time.time() - os.path.getmtime(p)) > REBUILD_DAYS * 86400

def _ensure_fresh():
    if _stale():
        with _lock:
            if _stale():
                try:
                    build()
                except Exception as e:
                    print("rebuild failed:", e)

def serve():
    import http.server, socketserver
    os.chdir(HERE or ".")
    class H(http.server.SimpleHTTPRequestHandler):
        def do_GET(self):
            if self.path in ("/", "/index.html"):
                _ensure_fresh()
            return http.server.SimpleHTTPRequestHandler.do_GET(self)
    port = int(os.environ.get("PORT", 8080))
    with socketserver.TCPServer(("0.0.0.0", port), H) as s:
        print("serving on port", port, "| auto-rebuild when older than", REBUILD_DAYS, "days")
        s.serve_forever()

if __name__ == "__main__":
    build()
    if len(sys.argv) < 2 or sys.argv[1] != "build":
        serve()
